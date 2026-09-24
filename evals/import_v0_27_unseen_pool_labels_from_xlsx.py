"""Import completed human labels from the blind workbook into the canonical CSV.

This script NEVER runs the judge. It requires all 60 rows to be LABELLED with a
0..4 score and nonblank human reason before replacing the CSV human-label fields.

Recommended invocation:

    uv run --with openpyxl python evals/import_v0_27_unseen_pool_labels_from_xlsx.py
"""
from __future__ import annotations

import argparse
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parents[1]
DATASETS = REPO_ROOT / "evals" / "datasets"
DEFAULT_CSV = DATASETS / "semantic_relevance_unseen_pool.v3.0.0.csv"
DEFAULT_XLSX = DATASETS / "semantic_relevance_unseen_pool.v3.0.0_labeling.xlsx"
BACKUP = DATASETS / "semantic_relevance_unseen_pool.v3.0.0.pre_label_import.csv"
EXPECTED_COLUMNS = [
    "case_id", "query_id", "query", "title", "authors", "description",
    "human_score", "human_reason", "review_status",
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Import completed unseen-pool human labels from Excel.")
    p.add_argument("--csv", type=Path, default=DEFAULT_CSV)
    p.add_argument("--workbook", type=Path, default=DEFAULT_XLSX)
    return p.parse_args()


def resolve(path: Path) -> Path:
    return path if path.is_absolute() else REPO_ROOT / path


def main() -> None:
    args = parse_args()
    csv_path = resolve(args.csv)
    xlsx_path = resolve(args.workbook)
    if not csv_path.exists():
        raise FileNotFoundError(csv_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(xlsx_path)

    source = pd.read_csv(csv_path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    if len(source) != 60 or source["case_id"].duplicated().any():
        raise ValueError("Canonical unseen-pool CSV must contain 60 unique case IDs.")

    wb = load_workbook(xlsx_path, data_only=False, read_only=True)
    if "Labeling" not in wb.sheetnames:
        raise ValueError("Workbook is missing the Labeling sheet.")
    ws = wb["Labeling"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Labeling sheet is empty.")
    headers = [str(x or "").strip() for x in rows[0]]
    if headers != EXPECTED_COLUMNS:
        raise ValueError(
            "Labeling sheet headers changed. Expected exactly: " + ", ".join(EXPECTED_COLUMNS)
        )

    data = pd.DataFrame(rows[1:], columns=headers).fillna("")
    data = data.loc[data["case_id"].astype(str).str.strip().ne("")].copy()
    if len(data) != 60 or data["case_id"].duplicated().any():
        raise ValueError("Labeling sheet must contain exactly 60 unique case IDs.")
    if set(data["case_id"].astype(str)) != set(source["case_id"].astype(str)):
        raise ValueError("Workbook case IDs do not exactly match canonical pool CSV.")

    # Evidence identity must not have been edited while labeling.
    compare_cols = ["query_id", "query", "title", "authors", "description"]
    joined = source[["case_id"] + compare_cols].merge(
        data[["case_id"] + compare_cols], on="case_id", suffixes=("_csv", "_xlsx"), validate="one_to_one"
    )
    for col in compare_cols:
        mismatch = joined[f"{col}_csv"].astype(str) != joined[f"{col}_xlsx"].astype(str)
        if mismatch.any():
            bad = joined.loc[mismatch, "case_id"].tolist()
            raise ValueError(f"Workbook evidence column {col} was modified for cases: {bad[:10]}")

    scores = pd.to_numeric(data["human_score"], errors="coerce")
    if scores.isna().any() or not scores.isin([0, 1, 2, 3, 4]).all():
        raise ValueError("All 60 human_score values must be complete integers 0..4.")
    if data["human_reason"].astype(str).str.strip().eq("").any():
        raise ValueError("All 60 human_reason values must be nonblank.")
    if not data["review_status"].astype(str).str.upper().eq("LABELLED").all():
        raise ValueError("All 60 review_status values must be LABELLED before import.")

    labels = data[["case_id", "human_score", "human_reason", "review_status"]].copy()
    labels["human_score"] = scores.astype(int).astype(str)
    labels["review_status"] = "LABELLED"

    updated = source.drop(columns=["human_score", "human_reason", "review_status"]).merge(
        labels, on="case_id", how="left", validate="one_to_one"
    )
    # Restore original canonical row order.
    order = {cid: i for i, cid in enumerate(source["case_id"])}
    updated["_order"] = updated["case_id"].map(order)
    updated = updated.sort_values("_order").drop(columns="_order")

    backup_path = resolve(BACKUP)
    if backup_path.exists():
        raise FileExistsError(
            f"Refusing to overwrite existing pre-import backup: {backup_path}. "
            "Preserve it as audit evidence."
        )
    shutil.copy2(csv_path, backup_path)
    updated.to_csv(csv_path, index=False, encoding="utf-8")

    print("v0.27 unseen-pool labels imported")
    print("---------------------------------")
    print("PASS  60/60 workbook rows matched canonical evidence")
    print("PASS  60/60 scores are 0..4")
    print("PASS  60/60 human reasons present")
    print("PASS  60/60 review_status=LABELLED")
    print(f"Backup: {backup_path}")
    print(f"Updated CSV: {csv_path}")
    print("NEXT: uv run python evals/validate_v0_27_unseen_pool.py")


if __name__ == "__main__":
    main()
