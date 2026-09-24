"""Export the populated v0.27 unseen pool to a blind human-labeling workbook.

This script does not run the judge. It intentionally excludes retrieval/source
metadata from the visible reviewer sheet to reduce anchoring bias.

Recommended invocation from PowerShell:

    uv run --with openpyxl python evals/export_v0_27_unseen_pool_labeling_xlsx.py
"""
from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO_ROOT = Path(__file__).resolve().parents[1]
DATASETS = REPO_ROOT / "evals" / "datasets"
DEFAULT_INPUT = DATASETS / "semantic_relevance_unseen_pool.v3.0.0.csv"
DEFAULT_OUTPUT = DATASETS / "semantic_relevance_unseen_pool.v3.0.0_labeling.xlsx"

VISIBLE_COLUMNS = [
    "case_id",
    "query_id",
    "query",
    "title",
    "authors",
    "description",
    "human_score",
    "human_reason",
    "review_status",
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Create blind-labeling workbook from populated unseen-pool CSV.")
    p.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    p.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return p.parse_args()


def resolve(path: Path) -> Path:
    return path if path.is_absolute() else REPO_ROOT / path


def validate_source(df: pd.DataFrame) -> None:
    required = set(VISIBLE_COLUMNS) | {
        "candidate_source", "retrieval_rank", "isbn13", "dataset_version", "rubric_version"
    }
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Pool is missing columns: {sorted(missing)}")
    if len(df) != 60:
        raise ValueError(f"Expected 60 pool rows; found {len(df)}.")
    for col in ("case_id", "query_id", "query", "title", "authors", "description"):
        if df[col].astype(str).str.strip().eq("").any():
            raise ValueError(f"Cannot export labeling workbook: {col} contains blanks.")
    if df["case_id"].duplicated().any():
        raise ValueError("Cannot export labeling workbook with duplicate case IDs.")


def main() -> None:
    args = parse_args()
    input_path = resolve(args.input)
    output_path = resolve(args.output)
    if not input_path.exists():
        raise FileNotFoundError(f"Populated unseen-pool CSV not found: {input_path}")

    df = pd.read_csv(input_path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    validate_source(df)

    wb = Workbook()
    ws = wb.active
    ws.title = "Labeling"
    guide = wb.create_sheet("Instructions")
    meta = wb.create_sheet("Metadata")

    # ------------------------------------------------------------------
    # Reviewer sheet: query + book evidence + human annotation only.
    # No retrieval rank/source is shown here.
    # ------------------------------------------------------------------
    ws.append(VISIBLE_COLUMNS)
    for _, row in df.iterrows():
        ws.append([row[c] for c in VISIBLE_COLUMNS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    evidence_fill = PatternFill("solid", fgColor="EAF2F8")
    bad_fill = PatternFill("solid", fgColor="F4CCCC")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2, max_row=61):
        for idx in range(1, 7):
            row[idx - 1].fill = evidence_fill
        for idx in range(7, 10):
            row[idx - 1].fill = editable_fill
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    score_validation = DataValidation(type="list", formula1='"0,1,2,3,4"', allow_blank=True)
    status_validation = DataValidation(type="list", formula1='"UNREVIEWED,LABELLED"', allow_blank=False)
    ws.add_data_validation(score_validation)
    ws.add_data_validation(status_validation)
    score_validation.add("G2:G61")
    status_validation.add("I2:I61")

    # Highlight a row marked LABELLED when score or reason is still missing.
    ws.conditional_formatting.add(
        "A2:I61",
        FormulaRule(formula=['AND($I2="LABELLED",OR($G2="",$H2=""))'], fill=bad_fill),
    )

    widths = {
        "A": 16, "B": 9, "C": 44, "D": 34, "E": 28,
        "F": 90, "G": 13, "H": 55, "I": 16,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for r in range(2, 62):
        ws.row_dimensions[r].height = 72
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = "A1:I61"

    # ------------------------------------------------------------------
    # Instructions/progress.
    # ------------------------------------------------------------------
    instructions = [
        ("v0.27 unseen pool — blind human labeling", ""),
        ("Purpose", "Human-label 60 fresh retrieval candidates before any judge output is generated."),
        ("What you judge", "Use only the query and supplied title/authors/description on the Labeling sheet."),
        ("human_score", "0=no relevance, 1=incidental/weak, 2=partial, 3=clear, 4=very strong/direct."),
        ("human_reason", "Briefly state which query concepts are supported or missing from the supplied description."),
        ("review_status", "Set to LABELLED only after score and reason are final."),
        ("Blindness", "Retrieval source/rank is intentionally hidden from the reviewer sheet."),
        ("Important", "Do not run the semantic relevance judge before all 60 rows are human-labelled and imported back to CSV."),
        ("", ""),
        ("Progress", ""),
        ("Total cases", 60),
        ("Labelled", '=COUNTIF(Labeling!I2:I61,"LABELLED")'),
        ("Scores entered", '=COUNT(Labeling!G2:G61)'),
        ("Completion %", '=B12/B11'),
    ]
    for row in instructions:
        guide.append(list(row))
    guide["A1"].font = Font(bold=True, size=14)
    guide["A1"].fill = header_fill
    guide["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 100
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide["B14"].number_format = "0%"

    # ------------------------------------------------------------------
    # Full provenance sheet: intentionally hidden during review.
    # ------------------------------------------------------------------
    meta.append(list(df.columns))
    for _, row in df.iterrows():
        meta.append([row[c] for c in df.columns])
    for cell in meta[1]:
        cell.fill = header_fill
        cell.font = header_font
    meta.freeze_panes = "A2"
    meta.sheet_state = "hidden"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    print("v0.27 blind-labeling workbook exported")
    print("---------------------------------------")
    print("PASS  60 populated evidence rows")
    print("PASS  retrieval source/rank hidden from reviewer sheet")
    print("PASS  score/status dropdown validation added")
    print("PASS  provenance preserved on hidden Metadata sheet")
    print(f"Workbook: {output_path}")


if __name__ == "__main__":
    main()
